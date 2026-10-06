"""Audit the pinned public TLCG table without training or saving source records."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
from importlib.metadata import version
import io
import json
import math
from pathlib import Path
import platform
import sys
from urllib.request import Request, urlopen
import uuid

from openpyxl import load_workbook

SOURCE_URL = (
    "https://data.mendeley.com/public-files/datasets/kdjvjbd9t7/"
    "files/3fced5f0-647d-4afa-8cde-d8e2942580f4/file_downloaded"
)
SOURCE_SHA256 = "76a2a6083d34125fea1c3e74ec19ec71a8f0df7fbc8bea6d34abc44cdf36a35d"
MAX_BYTES = 2_000_000


def audit(payload: bytes) -> dict:
    digest = hashlib.sha256(payload).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError("Source hash differs from audited release; review the new source explicitly.")
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        sheet = workbook["Foglio2"]
        iterator = sheet.iter_rows(values_only=True)
        next(iterator)  # Generic Column* header; semantic names are on row 2.
        headers = list(next(iterator))
        required = {"SUB", "TRIAL NUMBER", "LIE.TRUTH", "LIE.TRUTH_RT"}
        if len(set(headers)) != len(headers) or not required.issubset(headers):
            raise ValueError("Unexpected or ambiguous schema.")
        rows = [dict(zip(headers, values)) for values in iterator if any(v is not None for v in values)]
    finally:
        workbook.close()
    if not rows:
        raise ValueError("Empty source table.")
    for row in rows:
        if not row["SUB"] or not row["TRIAL NUMBER"] or row["LIE.TRUTH"] not in {"0", "1"}:
            raise ValueError("Missing identifiers or unexpected numeric label; no fallback is permitted.")
    groups = defaultdict(list)
    for row in rows:
        groups[row["SUB"]].append(row)
    keys = Counter((row["SUB"], row["TRIAL NUMBER"]) for row in rows)
    numeric_columns = headers[10:] + ["TRIAL NUMBER", "LIE.TRUTH_RT"]
    numeric_summary = {}
    for column in numeric_columns:
        values = []
        missing = 0
        invalid = 0
        for row in rows:
            value = row[column]
            if value is None or str(value).strip().upper() in {"", "NA", "NAN"}:
                missing += 1
                continue
            try:
                number = float(str(value).replace(",", "."))
            except ValueError:
                invalid += 1
                continue
            if not math.isfinite(number):
                invalid += 1
                continue
            values.append(number)
        numeric_summary[column] = {
            "missing_including_NA": missing,
            "invalid_or_nonfinite": invalid,
            "zero_count": sum(value == 0 for value in values),
            "negative_count": sum(value < 0 for value in values),
            "min": min(values, default=None),
            "max": max(values, default=None),
        }
    return {
        "provenance": {
            "source_url": SOURCE_URL,
            "dataset": "https://doi.org/10.17632/kdjvjbd9t7.2",
            "sha256": digest,
            "bytes": len(payload),
            "license": "CC BY 4.0",
            "checked_at_utc": datetime.now(timezone.utc).isoformat(),
            "python": sys.version.split()[0],
            "openpyxl": version("openpyxl"),
            "platform": platform.platform(),
            "raw_records_saved": False,
            "seed": None,
            "split_manifest": None,
            "training_performed": False,
        },
        "rows": len(rows),
        "columns": headers,
        "participants": len(groups),
        "numeric_label_counts": dict(Counter(row["LIE.TRUTH"] for row in rows)),
        "label_semantics": "UNVERIFIED: do not assign truth/lie meanings to numeric codes",
        "participants_with_both_codes": sum(
            {row["LIE.TRUTH"] for row in group} == {"0", "1"} for group in groups.values()
        ),
        "trial_count_distribution": dict(sorted(Counter(len(group) for group in groups.values()).items())),
        "duplicate_participant_trial_keys": sum(count - 1 for count in keys.values()),
        "duplicate_full_rows": len(rows) - len({tuple(row.values()) for row in rows}),
        "numeric_columns": numeric_summary,
        "any_missing_numeric_rows": sum(
            any(row[column] is None or str(row[column]).strip().upper() in {"", "NA", "NAN"}
                for column in numeric_columns)
            for row in rows
        ),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, help="Optional local trusted XLSX with the pinned source hash.")
    parser.add_argument("--output-dir", type=Path, help="New directory; existing directories are refused.")
    args = parser.parse_args()
    if args.source:
        if args.source.stat().st_size > MAX_BYTES:
            raise ValueError("Source exceeds expected size limit.")
        payload = args.source.read_bytes()
    else:
        request = Request(SOURCE_URL, headers={"User-Agent": "research-dataset-audit"})
        with urlopen(request, timeout=30) as response:
            payload = response.read(MAX_BYTES + 1)
        if len(payload) > MAX_BYTES:
            raise ValueError("Download exceeds expected size limit.")
    result = audit(payload)
    output_dir = args.output_dir or Path("outputs") / f"mendeley-audit-{uuid.uuid4().hex[:12]}"
    output_dir.mkdir(parents=True, exist_ok=False)
    output = output_dir / "audit_summary.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Summary: {output.resolve()}")
    print(f"Rows: {result['rows']}; participants: {result['participants']}; label meanings unverified.")


if __name__ == "__main__":
    main()
