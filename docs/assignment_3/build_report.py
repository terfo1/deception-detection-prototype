"""Build the editable report with bundled python-docx; not an analysis dependency."""
from __future__ import annotations

import argparse
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def append_text(paragraph, text: str, source_dir: Path, output_dir: Path):
    for part in re.split(r"(\[[^\]]+\]\([^)]+\))", text):
        match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if not match:
            paragraph.add_run(part)
            continue
        label, target = match.groups()
        if not target.startswith(("https://", "http://")):
            target = Path(os.path.relpath((source_dir / target).resolve(), output_dir.resolve())).as_posix()
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(
            target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        properties.append(fonts)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "24")
        properties.append(size)
        run.append(properties)
        value = OxmlElement("w:t")
        value.text = label
        run.append(value)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    source = Path(__file__).with_name("ASSIGNMENT_3_REPORT_EN.md").read_text(encoding="utf-8")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.9)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in list(fonts.attrib):
            if attribute.lower().endswith("theme"):
                del fonts.attrib[attribute]
        for name_key in ["ascii", "hAnsi", "eastAsia", "cs"]:
            fonts.set(qn(f"w:{name_key}"), "Times New Roman")
        paragraph_properties = style.element.get_or_add_pPr()
        for border in list(paragraph_properties.findall(qn("w:pBdr"))):
            paragraph_properties.remove(border)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(6)
    document.styles["Heading 1"].font.bold = True
    document.styles["Heading 1"].paragraph_format.space_before = Pt(0)
    document.styles["Title"].font.bold = True
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = "Assignment 3 - Software Development and Integration"
    for page_index, page in enumerate(source.split("<!-- page -->")):
        for block in page.strip().split("\n\n"):
            if block.startswith("# "):
                paragraph = document.add_paragraph(block[2:], "Title")
            elif block.startswith("## "):
                paragraph = document.add_paragraph(block[3:], "Heading 1")
                paragraph.paragraph_format.page_break_before = page_index > 0
            else:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                append_text(paragraph, block.replace("\n", " "), Path(__file__).parent, args.output.parent)
            paragraph.paragraph_format.widow_control = True
            paragraph.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.output.exists():
        raise FileExistsError(args.output)
    document.save(args.output)
    print(f"Report saved: {args.output}; {len(source.split())} source words; 10 planned pages")


if __name__ == "__main__":
    main()
